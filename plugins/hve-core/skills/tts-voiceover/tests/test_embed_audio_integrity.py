# Copyright (c) 2026 Microsoft Corporation. All rights reserved.
# SPDX-License-Identifier: MIT
"""Slide-to-WAV mapping, WAV validation, deck mutation, and save tests for embedding.

Decks and WAV files are synthetic and built at test time by ``conftest.py``.
"""

from __future__ import annotations

import logging

import embed_audio
from conftest import make_deck, make_malformed_wav, make_wav
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn


def _run(input_path, audio_dir, output_path):
    args = embed_audio.create_parser().parse_args(
        [
            "--input",
            str(input_path),
            "--audio-dir",
            str(audio_dir),
            "--output",
            str(output_path),
        ]
    )
    return embed_audio._run(args)


def _shape_counts(deck_path):
    return [len(slide.shapes) for slide in Presentation(str(deck_path)).slides]


class TestSlideMapping:
    """WAV files map to slides unambiguously before the deck is changed."""

    def test_given_duplicate_slide_wavs_when_embed_then_exits_2_without_output(
        self, tmp_path
    ):
        # Arrange
        deck = make_deck(tmp_path / "deck.pptx")
        make_wav(tmp_path / "audio" / "slide-1.wav")
        make_wav(tmp_path / "audio" / "slide-001.wav")
        output = tmp_path / "narrated.pptx"

        # Act
        exit_code = _run(deck, tmp_path / "audio", output)

        # Assert
        assert exit_code == 2
        assert not output.exists()

    def test_given_missing_wav_when_embed_then_slide_skipped(self, tmp_path):
        # Arrange
        deck = make_deck(tmp_path / "deck.pptx", slide_count=2)
        make_wav(tmp_path / "audio" / "slide-002.wav")
        output = tmp_path / "narrated.pptx"

        # Act
        exit_code = _run(deck, tmp_path / "audio", output)

        # Assert
        assert exit_code == 0
        assert _shape_counts(output) == [0, 1]

    def test_given_wav_beyond_deck_length_when_embed_then_warns_and_ignores_it(
        self, tmp_path, caplog
    ):
        # Arrange
        deck = make_deck(tmp_path / "deck.pptx", slide_count=2)
        make_wav(tmp_path / "audio" / "slide-001.wav")
        make_wav(tmp_path / "audio" / "slide-005.wav")
        output = tmp_path / "narrated.pptx"

        # Act
        with caplog.at_level(logging.WARNING, logger="embed_audio"):
            exit_code = _run(deck, tmp_path / "audio", output)

        # Assert
        assert exit_code == 0
        assert "slide-005.wav" in caplog.text
        assert _shape_counts(output) == [1, 0]


class TestWavValidation:
    """An unreadable WAV adds nothing to its slide; other slides still embed."""

    def test_given_malformed_wav_when_embed_then_no_orphan_shape_and_exit_1(
        self, tmp_path
    ):
        # Arrange
        deck = make_deck(tmp_path / "deck.pptx", slide_count=2)
        make_malformed_wav(tmp_path / "audio" / "slide-001.wav")
        make_wav(tmp_path / "audio" / "slide-002.wav")
        output = tmp_path / "narrated.pptx"

        # Act
        exit_code = _run(deck, tmp_path / "audio", output)

        # Assert
        assert exit_code == 1
        assert _shape_counts(output) == [0, 1]


class TestDeckMutation:
    """Existing timing and transitions are replaced (current behavior)."""

    def test_given_existing_timing_and_transition_when_embed_then_both_replaced(
        self, tmp_path, caplog
    ):
        # Arrange
        deck = make_deck(tmp_path / "deck.pptx", slide_count=1, existing_timing_on=(1,))
        make_wav(tmp_path / "audio" / "slide-001.wav")
        output = tmp_path / "narrated.pptx"

        # Act
        with caplog.at_level(logging.WARNING, logger="embed_audio"):
            exit_code = _run(deck, tmp_path / "audio", output)

        # Assert
        slide_element = Presentation(str(output)).slides[0]._element
        timings = slide_element.findall(qn("p:timing"))
        transitions = slide_element.findall(qn("p:transition"))
        assert exit_code == 0
        assert "Replacing existing slide timing" in caplog.text
        assert len(timings) == 1
        assert "playFrom(0)" in etree.tostring(timings[0], encoding="unicode")
        assert [t.get("advClick") for t in transitions] == ["0"]


class TestOutputHandling:
    """Output path, save failure, and empty-embedding behavior."""

    def test_given_existing_output_deck_when_embed_then_overwritten(self, tmp_path):
        # Arrange
        deck = make_deck(tmp_path / "deck.pptx")
        make_wav(tmp_path / "audio" / "slide-001.wav")
        output = tmp_path / "narrated.pptx"
        output.write_bytes(b"previous")

        # Act
        exit_code = _run(deck, tmp_path / "audio", output)

        # Assert
        assert exit_code == 0
        assert _shape_counts(output) == [1, 0]

    def test_given_no_wavs_when_embed_then_exits_1_without_output(self, tmp_path):
        # Arrange
        deck = make_deck(tmp_path / "deck.pptx")
        (tmp_path / "audio").mkdir()
        output = tmp_path / "narrated.pptx"

        # Act
        exit_code = _run(deck, tmp_path / "audio", output)

        # Assert
        assert exit_code == 1
        assert not output.exists()

    def test_given_output_equal_to_input_when_embed_then_exits_2_and_input_unchanged(
        self, tmp_path
    ):
        # Arrange
        deck = make_deck(tmp_path / "deck.pptx")
        original = deck.read_bytes()
        make_wav(tmp_path / "audio" / "slide-001.wav")

        # Act
        exit_code = _run(deck, tmp_path / "audio", deck)

        # Assert
        assert exit_code == 2
        assert deck.read_bytes() == original

    def test_given_save_failure_when_embed_then_exits_1(self, tmp_path, mocker):
        # Arrange
        deck = make_deck(tmp_path / "deck.pptx")
        make_wav(tmp_path / "audio" / "slide-001.wav")
        output = tmp_path / "narrated.pptx"
        mocker.patch(
            "pptx.presentation.Presentation.save", side_effect=OSError("disk full")
        )

        # Act
        exit_code = _run(deck, tmp_path / "audio", output)

        # Assert
        assert exit_code == 1
        assert not output.exists()
