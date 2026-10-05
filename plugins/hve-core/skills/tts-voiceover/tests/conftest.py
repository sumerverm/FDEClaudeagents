# Copyright (c) 2026 Microsoft Corporation. All rights reserved.
# SPDX-License-Identifier: MIT
"""Shared fixtures for tts-voiceover tests.

Provides an in-process Azure Speech SDK stand-in, a fake Entra credential, an
autouse network-deny guard, credential environment isolation, and builders
for synthetic WAV files, slide content trees, and PPTX decks. Nothing here
contacts Azure or reads real speaker notes.
"""

from __future__ import annotations

import datetime
import os
import socket
import sys
import types
import wave
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

PLACEHOLDER_KEY = "placeholder-speech-key"
PLACEHOLDER_RESOURCE_ID = (
    "/subscriptions/00000000-0000-0000-0000-000000000000/resourceGroups/rg"
    "/providers/Microsoft.CognitiveServices/accounts/placeholder"
)

_SPEECH_ENV_VARS = ("SPEECH_KEY", "SPEECH_RESOURCE_ID", "SPEECH_REGION")


@pytest.fixture(autouse=True)
def deny_network(monkeypatch):
    """Fail any test that attempts to open a network connection."""

    def _blocked(*_args, **_kwargs):
        raise RuntimeError("Network access is blocked in tts-voiceover tests")

    monkeypatch.setattr(socket.socket, "connect", _blocked)
    monkeypatch.setattr(socket.socket, "connect_ex", _blocked)
    monkeypatch.setattr(socket, "create_connection", _blocked)


@pytest.fixture(autouse=True)
def isolate_speech_environment(monkeypatch):
    """Remove Speech and Azure credential variables inherited from the host."""
    for name in _SPEECH_ENV_VARS:
        monkeypatch.delenv(name, raising=False)
    for name in list(os.environ):
        if name.startswith("AZURE_"):
            monkeypatch.delenv(name, raising=False)


class FakeSpeechConfig:
    def __init__(self, subscription=None, region=None, auth_token=None):
        self.subscription = subscription
        self.region = region
        self.auth_token = auth_token
        self.output_format = None

    def set_speech_synthesis_output_format(self, output_format):
        self.output_format = output_format


class FakeAudioOutputConfig:
    def __init__(self, filename):
        self.filename = filename


class FakeSpeechSdk:
    """Programmable stand-in for ``azure.cognitiveservices.speech``.

    ``outcomes`` is consumed one entry per synthesis call; when empty, calls
    succeed. Supported outcomes: ``"success"``, ``"cancel"``, and
    ``"cancel-partial"`` (writes bytes to the WAV path, then cancels).
    """

    def __init__(self) -> None:
        self.configs: list[FakeSpeechConfig] = []
        self.synthesis_calls: list[dict] = []
        self.outcomes: list[str] = []
        self.module = self._build_module()

    def _build_module(self) -> types.ModuleType:
        sdk = self

        class _SpeechConfig(FakeSpeechConfig):
            def __init__(self, *args, **kwargs):
                super().__init__(*args, **kwargs)
                sdk.configs.append(self)

        class _Synthesizer:
            def __init__(self, speech_config, audio_config):
                self.speech_config = speech_config
                self.audio_config = audio_config

            def speak_ssml_async(self, ssml):
                return _Future(sdk, self, ssml)

        module = types.ModuleType("azure.cognitiveservices.speech")
        module.SpeechConfig = _SpeechConfig
        module.SpeechSynthesizer = _Synthesizer
        module.ResultReason = types.SimpleNamespace(
            SynthesizingAudioCompleted="completed", Canceled="canceled"
        )
        module.SpeechSynthesisOutputFormat = types.SimpleNamespace(
            Riff24Khz16BitMonoPcm="riff-24khz-16bit-mono-pcm"
        )
        module.audio = types.SimpleNamespace(AudioOutputConfig=FakeAudioOutputConfig)
        return module


class _Future:
    def __init__(self, sdk: FakeSpeechSdk, synthesizer, ssml: str) -> None:
        self._sdk = sdk
        self._synthesizer = synthesizer
        self._ssml = ssml

    def get(self):
        outcome = self._sdk.outcomes.pop(0) if self._sdk.outcomes else "success"
        wav_path = Path(self._synthesizer.audio_config.filename)
        self._sdk.synthesis_calls.append(
            {
                "ssml": self._ssml,
                "config": self._synthesizer.speech_config,
                "wav_path": wav_path,
            }
        )
        if outcome == "success":
            make_wav(wav_path, duration_ms=100)
            return types.SimpleNamespace(
                reason="completed",
                audio_duration=datetime.timedelta(milliseconds=100),
            )
        if outcome == "cancel-partial":
            wav_path.write_bytes(b"RIFF")
        return types.SimpleNamespace(
            reason="canceled",
            cancellation_details=types.SimpleNamespace(
                reason="Error", error_details="synthetic cancellation"
            ),
        )


@pytest.fixture()
def fake_speech(monkeypatch) -> FakeSpeechSdk:
    """Install the fake SDK under the import name used by the scripts."""
    import azure.cognitiveservices

    sdk = FakeSpeechSdk()
    monkeypatch.setitem(sys.modules, "azure.cognitiveservices.speech", sdk.module)
    monkeypatch.setattr(azure.cognitiveservices, "speech", sdk.module, raising=False)
    return sdk


class FakeCredential:
    """Stand-in for ``DefaultAzureCredential`` with scripted token results."""

    def __init__(self, results: list) -> None:
        self._results = list(results)
        self.scopes: list[str] = []

    def get_token(self, scope):
        self.scopes.append(scope)
        result = self._results.pop(0)
        if isinstance(result, BaseException):
            raise result
        return result


@pytest.fixture()
def install_credential(monkeypatch):
    """Return a function that installs a ``FakeCredential`` for Entra auth."""
    import azure.identity

    def _install(results: list) -> FakeCredential:
        credential = FakeCredential(results)
        monkeypatch.setattr(
            azure.identity, "DefaultAzureCredential", lambda: credential
        )
        return credential

    return _install


def make_token(token: str, expires_on: float) -> types.SimpleNamespace:
    return types.SimpleNamespace(token=token, expires_on=int(expires_on))


def make_wav(path: Path, duration_ms: int = 100) -> Path:
    """Write a minimal silent mono 16 kHz WAV file."""
    sample_rate = 16000
    frames = int(sample_rate * duration_ms / 1000)
    path.parent.mkdir(parents=True, exist_ok=True)
    with wave.open(str(path), "wb") as handle:
        handle.setnchannels(1)
        handle.setsampwidth(2)
        handle.setframerate(sample_rate)
        handle.writeframes(b"\x00\x00" * frames)
    return path


def make_malformed_wav(path: Path) -> Path:
    """Write bytes that are not a readable RIFF/WAVE file."""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"not-a-wave-file")
    return path


def make_content_tree(root: Path, notes: list[str | None]) -> Path:
    """Write ``slide-NNN/content.yaml`` files with synthetic speaker notes."""
    import yaml

    for index, note in enumerate(notes, start=1):
        slide_dir = root / f"slide-{index:03d}"
        slide_dir.mkdir(parents=True, exist_ok=True)
        data = {"slide": index, "title": f"Synthetic {index}"}
        if note is not None:
            data["speaker_notes"] = note
        (slide_dir / "content.yaml").write_text(yaml.safe_dump(data), encoding="utf-8")
    return root


_EXISTING_TIMING = (
    '<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
    '<p:tnLst><p:par><p:cTn id="1" dur="indefinite" nodeType="tmRoot">'
    '<p:childTnLst><p:seq concurrent="1" nextAc="seek">'
    '<p:cTn id="2" dur="indefinite" nodeType="mainSeq"/>'
    "</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"
)


def make_deck(
    path: Path, slide_count: int = 2, *, existing_timing_on: tuple[int, ...] = ()
) -> Path:
    """Write a blank-layout deck; optionally add timing and a transition."""
    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]
    for index in range(1, slide_count + 1):
        slide = presentation.slides.add_slide(blank_layout)
        if index in existing_timing_on:
            slide._element.append(slide._element.makeelement(qn("p:transition"), {}))
            parser = etree.XMLParser(resolve_entities=False, no_network=True)
            slide._element.append(etree.fromstring(_EXISTING_TIMING, parser))
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(path))
    return path
